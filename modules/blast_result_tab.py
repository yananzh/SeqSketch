from PyQt6.QtWidgets import (
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QTableWidget,
    QTableWidgetItem,
    QPushButton,
    QFileDialog,
    QLabel,
    QMessageBox,
    QFrame,
    QDoubleSpinBox,
    QLineEdit,
    QSpinBox,
)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment as XlAlignment
from openpyxl.utils import get_column_letter

# outfmt 6 column definition  (12 standard fields)
_COLUMNS = [
    ("Query ID", "qseqid"),
    ("Subject ID", "sseqid"),
    ("Identity %", "pident"),
    ("Align Len", "length"),
    ("Mismatch", "mismatch"),
    ("Gap Open", "gapopen"),
    ("Q Start", "qstart"),
    ("Q End", "qend"),
    ("S Start", "sstart"),
    ("S End", "send"),
    ("E-value", "evalue"),
    ("Bit Score", "bitscore"),
]
_HDR = [c[0] for c in _COLUMNS]
_KEYS = [c[1] for c in _COLUMNS]

# Numeric columns (right-align)
_NUM_COLS = {2, 3, 4, 5, 6, 7, 8, 9, 10, 11}


class _NumericTableItem(QTableWidgetItem):
    """QTableWidgetItem that compares numerically when column is in _NUM_COLS."""

    def __lt__(self, other: QTableWidgetItem) -> bool:
        col = self.column()
        if col in _NUM_COLS:
            try:
                return float(self.text()) < float(other.text())
            except (ValueError, TypeError):
                pass  # fall through to text comparison
        return super().__lt__(other)


def _identity_color(pident_str: str) -> QColor | None:
    try:
        v = float(pident_str)
    except ValueError:
        return None
    if v >= 90:
        return QColor(200, 240, 200)  # green
    if v >= 60:
        return QColor(255, 245, 180)  # yellow
    return QColor(255, 210, 210)  # red


class BlastResultTab(QWidget):
    """Tab showing BLAST results from a TSV file (outfmt 6)."""

    def __init__(self, tsv_path: str, parent=None):
        super().__init__(parent)
        self.tsv_path = tsv_path
        self._rows: list[list[str]] = []
        self._build_ui()
        self._load_tsv(tsv_path)

    # ------------------------------------------------------------------ UI

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setSpacing(8)

        # ── header bar ────────────────────────────────────────────────────
        self._title_lbl = QLabel(
            f"<b>BLAST Results</b>  —  {os.path.basename(self.tsv_path)}"
        )
        self._title_lbl.setStyleSheet("color: #444; padding: 2px 0;")
        root.addWidget(self._title_lbl)
        root.addWidget(self._hline())

        filter_row = QHBoxLayout()
        filter_row.setSpacing(8)
        filter_row.addWidget(QLabel("Min identity %"))
        self.min_identity_spin = QDoubleSpinBox()
        self.min_identity_spin.setRange(0.0, 100.0)
        self.min_identity_spin.setDecimals(1)
        self.min_identity_spin.setSingleStep(5.0)
        self.min_identity_spin.valueChanged.connect(self._apply_filters)
        filter_row.addWidget(self.min_identity_spin)

        filter_row.addWidget(QLabel("Max e-value"))
        self.max_evalue_edit = QLineEdit()
        self.max_evalue_edit.setPlaceholderText("Any")
        self.max_evalue_edit.textChanged.connect(self._apply_filters)
        filter_row.addWidget(self.max_evalue_edit)

        filter_row.addWidget(QLabel("Min align len"))
        self.min_length_spin = QSpinBox()
        self.min_length_spin.setRange(0, 1000000)
        self.min_length_spin.valueChanged.connect(self._apply_filters)
        filter_row.addWidget(self.min_length_spin)

        self.reset_filters_btn = QPushButton("Reset Filters")
        self.reset_filters_btn.clicked.connect(self._reset_filters)
        filter_row.addWidget(self.reset_filters_btn)
        filter_row.addStretch()
        root.addLayout(filter_row)
        root.addWidget(self._hline())

        # ── table ────────────────────────────────────────────────────────
        self.table = QTableWidget()
        self.table.setColumnCount(len(_COLUMNS))
        self.table.setHorizontalHeaderLabels(_HDR)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.setSortingEnabled(True)
        self.table.horizontalHeader().setStretchLastSection(True)
        root.addWidget(self.table)

        # ── stat label ────────────────────────────────────────────────────
        self._stat_lbl = QLabel("")
        self._stat_lbl.setStyleSheet("color: #666; font-size: 11px;")
        root.addWidget(self._stat_lbl)

    @staticmethod
    def _hline():
        f = QFrame()
        f.setFrameShape(QFrame.Shape.HLine)
        f.setFrameShadow(QFrame.Shadow.Sunken)
        return f

    # ---------------------------------------------------------------- data

    def _load_tsv(self, path: str):
        if not path or not os.path.isfile(path):
            self._stat_lbl.setText(f"File not found: {path}")
            return

        self.tsv_path = path
        self._title_lbl.setText(f"<b>BLAST Results</b>  —  {os.path.basename(path)}")
        self._rows.clear()
        self.table.setSortingEnabled(False)
        self.table.setRowCount(0)

        try:
            with open(path, "r", encoding="utf-8") as fh:
                for raw in fh:
                    line = raw.rstrip("\n")
                    if not line:
                        continue
                    parts = line.split("\t")
                    # skip header row written by blast_run_dialog
                    if parts[0] == "qseqid":
                        continue
                    # pad / trim to 12 columns
                    while len(parts) < 12:
                        parts.append("")
                    self._rows.append(parts[:12])
        except Exception as exc:
            QMessageBox.critical(self, "File Error", f"Could not read TSV:\n{exc}")
            return

        self._apply_filters()

    def _populate_table(self, rows: list[list[str]]):
        self.table.setSortingEnabled(False)
        self.table.setRowCount(len(rows))
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                item = _NumericTableItem(val) if c in _NUM_COLS else QTableWidgetItem(val)
                if c in _NUM_COLS:
                    item.setTextAlignment(
                        Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter
                    )
                self.table.setItem(r, c, item)

        self.table.setSortingEnabled(True)
        self.table.resizeColumnsToContents()

    def _matches_filters(self, row: list[str]) -> bool:
        try:
            identity = float(row[2])
        except ValueError:
            identity = 0.0
        if identity < self.min_identity_spin.value():
            return False

        try:
            align_len = int(float(row[3]))
        except ValueError:
            align_len = 0
        if align_len < self.min_length_spin.value():
            return False

        max_evalue_text = self.max_evalue_edit.text().strip()
        if max_evalue_text:
            try:
                max_evalue = float(max_evalue_text)
                evalue = float(row[10])
            except ValueError:
                return False
            if evalue > max_evalue:
                return False

        return True

    def _apply_filters(self):
        visible_rows = [row for row in self._rows if self._matches_filters(row)]
        self._populate_table(visible_rows)

        total = len(self._rows)
        visible = len(visible_rows)
        if visible:
            self._stat_lbl.setText(
                f"Showing {visible} / {total} hit{'s' if total != 1 else ''}."
            )
            self.table.selectRow(0)
        else:
            if total:
                self._stat_lbl.setText(
                    f"Showing 0 / {total} hits. Adjust the filters to see more results."
                )
            else:
                self._stat_lbl.setText("No hits found in this file.")

    def _reset_filters(self):
        self.min_identity_spin.setValue(0.0)
        self.max_evalue_edit.clear()
        self.min_length_spin.setValue(0)

    # ---------------------------------------------------------------- slots

    def _open_other_tsv(self):
        f, _ = QFileDialog.getOpenFileName(
            self,
            "Open BLAST TSV result",
            "",
            "TSV files (*.tsv *.txt);;All Files (*)",
        )
        if f:
            self._load_tsv(f)

    def _export(self, fmt: str):
        if self.table.rowCount() == 0:
            QMessageBox.warning(self, "No Data", "There are no results to export.")
            return
        sep = "," if fmt == "csv" else "\t"
        f, _ = QFileDialog.getSaveFileName(
            self,
            f"Export as {fmt.upper()}",
            f"blast_result.{fmt}",
            f"{fmt.upper()} files (*.{fmt});;All Files (*)",
        )
        if not f:
            return
        try:
            with open(f, "w", encoding="utf-8", newline="") as fh:
                writer = csv.writer(fh, delimiter=sep)
                writer.writerow(_HDR)
                for r in range(self.table.rowCount()):
                    writer.writerow([
                        self.table.item(r, c).text() if self.table.item(r, c) else ""
                        for c in range(12)
                    ])
            QMessageBox.information(self, "Exported", f"Results saved to:\n{f}")
        except Exception as exc:
            QMessageBox.critical(self, "Export Error", str(exc))

    def _export_xlsx(self):
        if self.table.rowCount() == 0:
            QMessageBox.warning(self, "No Data", "There are no results to export.")
            return
        f, _ = QFileDialog.getSaveFileName(
            self,
            "Export as XLSX",
            "blast_result.xlsx",
            "Excel files (*.xlsx);;All Files (*)",
        )
        if not f:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "BLAST Results"

            header_fill = PatternFill(
                start_color="2563EB", end_color="2563EB", fill_type="solid"
            )
            header_font = Font(color="FFFFFF", bold=True, size=11)
            header_align = XlAlignment(horizontal="center", vertical="center")

            for c, col_name in enumerate(_HDR, 1):
                cell = ws.cell(row=1, column=c, value=col_name)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_align

            green_fill = PatternFill(
                start_color="C8F0C8", end_color="C8F0C8", fill_type="solid"
            )
            yellow_fill = PatternFill(
                start_color="FFF5B4", end_color="FFF5B4", fill_type="solid"
            )
            red_fill = PatternFill(
                start_color="FFD2D2", end_color="FFD2D2", fill_type="solid"
            )

            for r in range(self.table.rowCount()):
                for c in range(12):
                    val = self.table.item(r, c).text() if self.table.item(r, c) else ""
                    cell = ws.cell(row=r + 2, column=c + 1, value=val)
                    if c in _NUM_COLS:
                        cell.alignment = XlAlignment(horizontal="right")
                    if c == 2:
                        try:
                            v = float(val)
                            if v >= 90:
                                cell.fill = green_fill
                            elif v >= 60:
                                cell.fill = yellow_fill
                            else:
                                cell.fill = red_fill
                        except ValueError:
                            pass

            for c in range(1, 13):
                ws.column_dimensions[get_column_letter(c)].width = 16

            ws.auto_filter.ref = ws.dimensions
            wb.save(f)
            QMessageBox.information(self, "Exported", f"Results saved to:\n{f}")
        except Exception as exc:
            QMessageBox.critical(self, "Export Error", str(exc))
